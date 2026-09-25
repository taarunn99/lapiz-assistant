#!/usr/bin/env python3
"""
Lapiz Blue stock sync backfill (history).

Goal: make physical stock (shipments and returns) agree with accounting stock (invoices and
credit notes) for every sales order, exactly the way the live Deluge functions do it going forward.

Per sales order, per LIVE (non void) invoice:
    one package + one shipment mirroring that invoice's SO lines, dated to the invoice date,
    notes "LB-STOCKSYNC inv:<invoice_id> <invoice_number> mod:<last_modified> backfill"
Per credit note WITH item lines tied to such an invoice:
    one sales return + receive mirroring the credit note lines, dated to the credit note date,
    notes "LB-STOCKSYNC cn:<creditnote_id> <cn_number> mod:<last_modified> backfill"

The note format matches the live functions, so if an old invoice is edited after the backfill the
live function recognises the backfilled package as its mirror and rebuilds it correctly.

Modes
  dry   (default)  read only. Writes plan.json + Backfill_Plan.xlsx. Creates nothing.
  live             reads plan.json from the dry run, re-verifies each SO right before acting,
                   then creates packages, shipments, returns and receives. Resumable (done.json).

Environment
  ZOHO_CLIENT_ID, ZOHO_CLIENT_SECRET, ZOHO_REFRESH_TOKEN, ZOHO_ORG_ID_DXB
Optional
  MAX_CALLS   hard cap on API calls per run (default 8000)
  SO_LIMIT    only process the first N sales orders (testing)
  SO_NUMBERS  comma separated SO numbers to restrict to (testing on ZZ TEST)
  SKIP_CN     set to 1 to skip the credit note scan (faster dry run)
"""
import json, os, sys, time, datetime as dt
import requests

BASE_INV = "https://www.zohoapis.com/inventory/v1"
TOKEN_URL = "https://accounts.zoho.com/oauth/v2/token"
ORG = os.environ["ZOHO_ORG_ID_DXB"]
MAX_CALLS = int(os.environ.get("MAX_CALLS", "8000"))
SO_LIMIT = int(os.environ.get("SO_LIMIT", "0") or 0)
SO_NUMBERS = {s.strip() for s in os.environ.get("SO_NUMBERS", "").split(",") if s.strip()}
SKIP_CN = os.environ.get("SKIP_CN", "0") == "1"
MODE = sys.argv[1] if len(sys.argv) > 1 else "dry"
OUT = os.environ.get("OUT_DIR", "out"); os.makedirs(OUT, exist_ok=True)
MARK = "LB-STOCKSYNC"
DELIVERY_METHOD = "Al Quoz"

def log(*a):
    print(*a, flush=True)

# ---------- auth + throttled client ----------
class Zoho:
    def __init__(self):
        self.calls = 0; self.window = []; self.token = None; self.exp = 0
    def _auth(self):
        if time.time() < self.exp - 60: return
        r = requests.post(TOKEN_URL, data={
            "client_id": os.environ["ZOHO_CLIENT_ID"], "client_secret": os.environ["ZOHO_CLIENT_SECRET"],
            "grant_type": "refresh_token", "refresh_token": os.environ["ZOHO_REFRESH_TOKEN"]}, timeout=30)
        j = r.json()
        if "access_token" not in j: raise SystemExit(f"auth failed: {j}")
        self.token = j["access_token"]; self.exp = time.time() + int(j.get("expires_in", 3600))
    def req(self, method, path, **kw):
        if self.calls >= MAX_CALLS: raise SystemExit(f"MAX_CALLS {MAX_CALLS} reached, stopping cleanly")
        now = time.time(); self.window = [t for t in self.window if now - t < 60]
        if len(self.window) >= 80: time.sleep(60 - (now - self.window[0]) + 0.5)
        self._auth()
        params = kw.pop("params", {}); params["organization_id"] = ORG
        for attempt in range(4):
            r = requests.request(method, BASE_INV + path, params=params,
                                 headers={"Authorization": f"Zoho-oauthtoken {self.token}"}, timeout=60, **kw)
            self.calls += 1; self.window.append(time.time())
            if r.status_code == 429:
                txt = r.text.lower()
                if "day" in txt or "daily" in txt or "limit exceeded" in txt or attempt >= 2:
                    # a 429 that survives two 65s waits is the daily cap, whatever the message says. Stop and save.
                    raise SystemExit(f"Zoho API limit (429) after {self.calls} calls this run: {r.text[:120]}")
                log("rate limited, sleeping 65s"); time.sleep(65); continue
            if r.status_code >= 500:
                time.sleep(5 * (attempt + 1)); continue
            try: j = r.json()
            except Exception: j = {"code": -1, "message": r.text[:200]}
            return j
        # never hand a half-answer back to the planner: a missing package or invoice would turn into a duplicate shipment
        raise SystemExit(f"Zoho API kept failing on {method} {path} (last status {r.status_code}); stopping and saving progress")
    def get(self, path, **params): return self.req("GET", path, params=params)
    def post(self, path, body, **params): return self.req("POST", path, params=params, json=body)

Z = Zoho()
INV_CACHE = {}
def get_invoice(inv_id):
    if inv_id not in INV_CACHE:
        j = Z.get(f"/invoices/{inv_id}")
        INV_CACHE[inv_id] = j.get("invoice") if j.get("code") == 0 else None
    return INV_CACHE[inv_id]

# ---------- discovery ----------
def list_candidate_sos():
    """SOs whose list row says invoiced > shipped. One call per 200 SOs."""
    page, out = 1, []
    while True:
        j = Z.get("/salesorders", filter_by="Status.ToBePacked", per_page=200, page=page, sort_column="date")
        if j.get("code") != 0: raise SystemExit(f"list failed: {j}")
        for s in j.get("salesorders", []):
            if SO_NUMBERS and s["salesorder_number"] not in SO_NUMBERS: continue
            if float(s.get("quantity_invoiced", 0)) > float(s.get("quantity_shipped", 0)):
                out.append(s)
        if not j.get("page_context", {}).get("has_more_page"): break
        page += 1
        if SO_LIMIT and len(out) >= SO_LIMIT: break
    return out[:SO_LIMIT] if SO_LIMIT else out

def so_by_numbers():
    """Testing helper: SO_NUMBERS may point at SOs that are already shipped (for return planning)."""
    out = []
    for n in SO_NUMBERS:
        j = Z.get("/salesorders", salesorder_number=n)
        out += j.get("salesorders", [])
    return out

# ---------- planning for one SO ----------
def plan_so(so_id):
    """Returns (plan, flags). plan['packages'] = per live invoice; plan['returns'] filled later."""
    j = Z.get(f"/salesorders/{so_id}")
    if j.get("code") != 0: return None, [f"fetch failed: {j.get('message')}"]
    so = j["salesorder"]; flags = []
    if so.get("is_dropshipped") or so.get("is_drop_shipment"): flags.append("DROP SHIP, skip")
    if so.get("status") in ("void", "draft"): flags.append(f"status {so['status']}, skip")
    lines = {l["line_item_id"]: l for l in so["line_items"] if l.get("item_type") == "inventory"}
    for lid, l in lines.items():
        if l.get("batches") or l.get("serial_numbers"): flags.append(f"tracked item {l['name']}, needs human")
    if any("skip" in f or "human" in f for f in flags): return None, flags

    # packages that already exist on this SO: which invoices do they already mirror?
    existing_mirror_inv = set(); other_packed = {lid: 0.0 for lid in lines}
    for pk in so.get("packages", []):
        pj = Z.get(f"/packages/{pk['package_id']}")
        if pj.get("code") != 0:
            # cannot see what this package covers, so do not plan anything for this SO
            return None, flags + [f"package {pk.get('package_number', pk['package_id'])} could not be read ({pj.get('message')}), needs human"]
        p = pj["package"]; notes = p.get("notes", "") or ""
        if MARK in notes and "inv:" in notes:
            existing_mirror_inv.add(notes.split("inv:")[1].split()[0])
        else:
            for pl in p.get("line_items", []):
                lid = pl.get("so_line_item_id")
                if lid in other_packed: other_packed[lid] += float(pl.get("quantity", 0))

    # live invoices, oldest first. Void ones are ignored entirely (Zoho still counts them on the SO).
    invoices = sorted([i for i in so.get("invoices", []) if i.get("status") not in ("void", "draft")],
                      key=lambda i: (i["date"], i["invoice_number"]))
    room = {lid: float(l["quantity"]) - float(l.get("quantity_packed", 0)) for lid, l in lines.items()}
    hand_packed_left = dict(other_packed)   # hand made packages are assumed to cover the OLDEST invoices
    plan = []
    for inv in invoices:
        if inv["invoice_id"] in existing_mirror_inv: continue   # already mirrored by an earlier run or the live function
        iv = get_invoice(inv["invoice_id"])
        if iv is None: return None, flags + [f"invoice {inv['invoice_number']} could not be read, needs human"]
        pkg_lines = []
        for il in iv["line_items"]:
            lid = il.get("salesorder_item_id")
            if not lid or lid not in lines or il.get("item_type") != "inventory": continue
            q = float(il["quantity"])
            covered = min(q, hand_packed_left.get(lid, 0.0)); hand_packed_left[lid] -= covered
            take = min(q - covered, max(room[lid], 0.0))
            if take > 0:
                pkg_lines.append({"so_line_item_id": lid, "item": lines[lid]["name"], "quantity": take})
                room[lid] -= take
            if q - covered - take > 0.0001:
                flags.append(f"{inv['invoice_number']} {lines[lid]['name']}: {q - covered - take:g} could not be packed (no room on SO)")
        if pkg_lines:
            plan.append({"invoice_id": inv["invoice_id"], "invoice_number": inv["invoice_number"],
                         "invoice_date": inv["date"], "invoice_mod": iv.get("last_modified_time", ""),
                         "lines": pkg_lines})
    return {"salesorder_id": so_id, "salesorder_number": so["salesorder_number"],
            "customer": so["customer_name"], "packages": plan, "returns": []}, flags

# ---------- credit notes with item lines -> returns ----------
def scan_credit_notes(plans_by_so):
    """Walk all live credit notes. For item lines tied to an SO-linked invoice, add a return to that SO's plan
    (creating a plan entry if the SO was not in the package list)."""
    page, seen, rows_skipped = 1, 0, []
    while True:
        j = Z.get("/creditnotes", per_page=200, page=page, sort_column="date")
        if j.get("code") != 0: log("credit note list failed", j); break
        for c in j.get("creditnotes", []):
            if c.get("status") in ("void", "draft"): continue
            cj = Z.get(f"/creditnotes/{c['creditnote_id']}")
            if cj.get("code") != 0: continue
            cn = cj["creditnote"]; seen += 1
            if cn.get("salesreturn_id"): continue   # born from a sales return, Zoho already moved the stock
            want = {}   # invoice_id -> {invoice_item_id: qty}
            for cl in cn.get("line_items", []):
                if cl.get("item_type") != "inventory": continue
                iid = cl.get("invoice_id") or cn.get("invoice_id"); iil = cl.get("invoice_item_id")
                if not iid or not iil:
                    rows_skipped.append([cn["creditnote_number"], cn["customer_name"], cl.get("name"), cl.get("quantity"), "item line not tied to an invoice line"]); continue
                want.setdefault(iid, {}); want[iid][iil] = want[iid].get(iil, 0.0) + float(cl["quantity"])
            for iid, lines in want.items():
                iv = get_invoice(iid)
                if iv is None: rows_skipped.append([cn["creditnote_number"], cn["customer_name"], "", "", f"invoice {iid} fetch failed"]); continue
                so_ids = [s["salesorder_id"] for s in iv.get("salesorders", [])] or ([iv["salesorder_id"]] if iv.get("salesorder_id") else [])
                if not so_ids:
                    rows_skipped.append([cn["creditnote_number"], cn["customer_name"], "", sum(lines.values()), f"invoice {iv['invoice_number']} has no sales order"]); continue
                by_so_line = {}
                for il in iv["line_items"]:
                    if il["line_item_id"] in lines and il.get("salesorder_item_id"):
                        by_so_line[il["salesorder_item_id"]] = by_so_line.get(il["salesorder_item_id"], 0.0) + lines[il["line_item_id"]]
                if not by_so_line: continue
                so_id = so_ids[0]
                if so_id not in plans_by_so:
                    sj = Z.get(f"/salesorders/{so_id}")
                    if sj.get("code") != 0: continue
                    s = sj["salesorder"]
                    plans_by_so[so_id] = {"salesorder_id": so_id, "salesorder_number": s["salesorder_number"],
                                          "customer": s["customer_name"], "packages": [], "returns": [], "_so": s}
                p = plans_by_so[so_id]
                # skip if a mirror return for this credit note already exists
                s = p.get("_so") or Z.get(f"/salesorders/{so_id}").get("salesorder", {})
                already = False
                for sr in s.get("salesreturns", []):
                    rj = Z.get(f"/salesreturns/{sr['salesreturn_id']}")
                    if rj.get("code") == 0 and f"cn:{cn['creditnote_id']}" in (rj["salesreturn"].get("notes") or ""):
                        already = True
                if already: continue
                so_lines = {l["line_item_id"]: l for l in s.get("line_items", [])}
                ret_lines = []
                for lid, q in by_so_line.items():
                    sl = so_lines.get(lid, {})
                    left = q - float(sl.get("quantity_returned", 0))   # manual returns already on the SO count
                    if left > 0:
                        ret_lines.append({"salesorder_item_id": lid, "item_id": sl.get("item_id"), "item": sl.get("name", lid), "quantity": left,
                                          "returnable": sl.get("is_returnable", True)})
                if not ret_lines: continue
                if any(not l["returnable"] for l in ret_lines):
                    rows_skipped.append([cn["creditnote_number"], cn["customer_name"], ", ".join(l["item"] for l in ret_lines if not l["returnable"]), "", "item is marked not returnable in Zoho"]); continue
                p["returns"].append({"creditnote_id": cn["creditnote_id"], "creditnote_number": cn["creditnote_number"],
                                     "creditnote_date": cn["date"], "creditnote_mod": cn.get("last_modified_time", ""),
                                     "invoice_number": iv["invoice_number"], "lines": ret_lines})
        if not j.get("page_context", {}).get("has_more_page"): break
        page += 1
    log(f"credit notes scanned: {seen}")
    return rows_skipped

# ---------- execution for one SO ----------
def execute_so(p):
    fresh, flags = plan_so(p["salesorder_id"])
    res = []
    if fresh is None:
        return [f"{p['salesorder_number']}: skipped now ({'; '.join(flags)})"]
    for pk in fresh["packages"]:
        note = f"{MARK} inv:{pk['invoice_id']} {pk['invoice_number']} mod:{pk['invoice_mod']} backfill"
        body = {"date": pk["invoice_date"], "notes": note,
                "line_items": [{"so_line_item_id": l["so_line_item_id"], "quantity": l["quantity"]} for l in pk["lines"]]}
        r = Z.post("/packages", body, salesorder_id=p["salesorder_id"])
        if r.get("code") != 0:
            res.append(f"{p['salesorder_number']} pkg for {pk['invoice_number']} FAILED: {r.get('message')}"); continue
        pid = r["package"]["package_id"]
        s = Z.post("/shipmentorders", {"date": pk["invoice_date"], "delivery_method": DELIVERY_METHOD, "notes": note},
                   package_ids=pid, salesorder_id=p["salesorder_id"], send_notification="false")
        if s.get("code") != 0:
            res.append(f"{p['salesorder_number']} shipment for {pk['invoice_number']} FAILED after package {pid}: {s.get('message')}"); continue
        res.append(f"{p['salesorder_number']} OK {pk['invoice_number']} pkg {r['package'].get('package_number')} ship {s['shipmentorder'].get('shipment_number')}")
    # returns from the plan (credit notes do not change between dry and live often; re-check room)
    if p.get("returns"):
        sj = Z.get(f"/salesorders/{p['salesorder_id']}"); so = sj.get("salesorder", {})
        room = {l["line_item_id"]: float(l.get("quantity_shipped", 0)) - float(l.get("quantity_returned", 0)) for l in so.get("line_items", [])}
        done_cn = set()
        for sr in so.get("salesreturns", []):
            rj = Z.get(f"/salesreturns/{sr['salesreturn_id']}")
            n = (rj.get("salesreturn", {}).get("notes") or "")
            if "cn:" in n: done_cn.add(n.split("cn:")[1].split()[0])
        for rt in p["returns"]:
            if rt["creditnote_id"] in done_cn: res.append(f"{p['salesorder_number']} return for {rt['creditnote_number']} already exists"); continue
            note = f"{MARK} cn:{rt['creditnote_id']} {rt['creditnote_number']} mod:{rt['creditnote_mod']} backfill"
            lines = []
            for l in rt["lines"]:
                q = min(l["quantity"], max(room.get(l["salesorder_item_id"], 0.0), 0.0))
                if q < l["quantity"]: res.append(f"{p['salesorder_number']} {rt['creditnote_number']} {l['item']}: capped {l['quantity']:g} -> {q:g} (not enough shipped)")
                if q > 0: lines.append({"salesorder_item_id": l["salesorder_item_id"], "item_id": l.get("item_id"), "quantity": q}); room[l["salesorder_item_id"]] -= q
            if not lines: continue
            r = Z.post("/salesreturns", {"salesorder_id": p["salesorder_id"], "date": rt["creditnote_date"],
                                          "reason": f"Credit note {rt['creditnote_number']}", "notes": note, "line_items": lines})
            if r.get("code") != 0:
                res.append(f"{p['salesorder_number']} return for {rt['creditnote_number']} FAILED: {r.get('message')}"); continue
            sret = r["salesreturn"]
            rcv = {"salesreturn_id": sret["salesreturn_id"], "date": rt["creditnote_date"], "notes": note,
                   "line_items": [{"line_item_id": x["line_item_id"], "quantity": x["quantity"]} for x in sret.get("line_items", [])]}
            rr = Z.post("/salesreturnreceives", rcv, salesreturn_id=sret["salesreturn_id"])
            if rr.get("code") != 0:
                res.append(f"{p['salesorder_number']} return {sret.get('salesreturn_number')} created but receive FAILED: {rr.get('message')}"); continue
            res.append(f"{p['salesorder_number']} OK return {sret.get('salesreturn_number')} for {rt['creditnote_number']} received")
    return res or [f"{p['salesorder_number']}: nothing to do now"]

# ---------- main ----------
def main():
    if MODE == "dry":
        sos = so_by_numbers() if SO_NUMBERS else list_candidate_sos()
        # resume: previous dry run artifact (plan.json + state.json) is reused, examined SOs are not re-fetched
        plans_by_so, rows, skipped = {}, [], []
        state = {"examined": [], "skipped": [], "cn_done": False}
        if os.path.exists(f"{OUT}/plan.json") and os.path.exists(f"{OUT}/state.json") and not SO_NUMBERS:
            for p in json.load(open(f"{OUT}/plan.json")): plans_by_so[p["salesorder_id"]] = p
            state = json.load(open(f"{OUT}/state.json"))
            skipped = state.get("skipped", [])
            log(f"resuming: {len(state['examined'])} SOs already examined, {len(plans_by_so)} planned")
        examined = set(state["examined"])
        sos = [s for s in sos if s["salesorder_id"] not in examined]
        log(f"{len(sos)} sales orders to examine this run")
        partial = ""

        def checkpoint(msg):
            # save what we have so far, so a cancelled or killed run loses at most 100 sales orders of work
            state.update({"examined": sorted(examined), "skipped": skipped, "cn_skipped": state.get("cn_skipped", [])})
            json.dump(state, open(f"{OUT}/state.json", "w"))
            json.dump([p for p in plans_by_so.values() if p["packages"] or p["returns"]], open(f"{OUT}/plan.json", "w"), indent=1)
            ck_rows = [[p["salesorder_number"], p["customer"], pk["invoice_number"], pk["invoice_date"], l["item"], l["quantity"], ""]
                       for p in plans_by_so.values() for pk in p["packages"] for l in pk["lines"]]
            write_xlsx(ck_rows, [], skipped, state.get("cn_skipped", []), list(plans_by_so.values()), msg)

        for i, s in enumerate(sos, 1):
            try:
                p, flags = plan_so(s["salesorder_id"])
                examined.add(s["salesorder_id"])
            except (SystemExit, KeyboardInterrupt) as e:
                partial = f"PARTIAL: {e}. Planned {i-1} of {len(sos)} sales orders. Run again tomorrow to continue (already mirrored ones are skipped automatically)."
                log(partial); break
            if p and p["packages"]:
                plans_by_so[s["salesorder_id"]] = p
                for pk in p["packages"]:
                    for l in pk["lines"]:
                        rows.append([p["salesorder_number"], p["customer"], pk["invoice_number"], pk["invoice_date"],
                                     l["item"], l["quantity"], "; ".join(flags)])
            elif p and not p["packages"] and not flags:
                pass   # nothing left to ship on this SO
            else:
                skipped.append([s["salesorder_number"], s["customer_name"], "; ".join(flags) or "no packable lines"])
            if i % 100 == 0:
                log(f"  {i}/{len(sos)} planned, {Z.calls} calls used")
                checkpoint(f"PARTIAL (checkpoint at {i}/{len(sos)}): the run did not finish. Run again to continue.")
        cn_skipped = state.get("cn_skipped", [])
        if not SKIP_CN and not partial and not state.get("cn_done"):
            try:
                cn_skipped = scan_credit_notes(plans_by_so); state["cn_done"] = True
            except (SystemExit, KeyboardInterrupt) as e:
                partial = f"PARTIAL: {e} during credit note scan. Shipments plan is complete, returns plan is incomplete, run again."
                log(partial)
        # rebuild rows from the full plan (previous runs included)
        rows = []
        for p in plans_by_so.values():
            for pk in p["packages"]:
                for l in pk["lines"]:
                    rows.append([p["salesorder_number"], p["customer"], pk["invoice_number"], pk["invoice_date"], l["item"], l["quantity"], ""])
        state.update({"examined": sorted(examined), "skipped": skipped, "cn_skipped": cn_skipped})
        json.dump(state, open(f"{OUT}/state.json", "w"))
        ret_rows = []
        for p in plans_by_so.values():
            p.pop("_so", None)
            for rt in p["returns"]:
                for l in rt["lines"]:
                    ret_rows.append([p["salesorder_number"], p["customer"], rt["creditnote_number"], rt["creditnote_date"],
                                     rt["invoice_number"], l["item"], l["quantity"]])
        plans = [p for p in plans_by_so.values() if p["packages"] or p["returns"]]
        json.dump(plans, open(f"{OUT}/plan.json", "w"), indent=1)
        write_xlsx(rows, ret_rows, skipped, cn_skipped, plans, partial)
        log(f"plan: {len(plans)} SOs, {sum(len(p['packages']) for p in plans)} packages, {len(rows)} lines, "
            f"{sum(len(p['returns']) for p in plans)} returns, {len(skipped)} skipped. calls {Z.calls}")
    elif MODE == "live":
        plans = json.load(open(f"{OUT}/plan.json"))
        done_path = f"{OUT}/done.json"
        done = set(json.load(open(done_path))) if os.path.exists(done_path) else set()
        logf = open(f"{OUT}/live_log.txt", "a")
        for p in plans:
            if p["salesorder_id"] in done: continue
            try:
                for line in execute_so(p): log(line); logf.write(line + "\n")
            except SystemExit as e:
                log(e); break
            done.add(p["salesorder_id"]); json.dump(sorted(done), open(done_path, "w"))
            logf.flush()
        log(f"live run stopped. {len(done)}/{len(plans)} SOs processed. calls {Z.calls}")
    else:
        raise SystemExit("mode must be dry or live")

def write_xlsx(rows, ret_rows, skipped, cn_skipped, plans, partial=""):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    def hdr(ws, cols, widths):
        ws.append(cols)
        for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
        for col, w in zip("ABCDEFGH", widths): ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
    wb = Workbook()
    ws3 = wb.active; ws3.title = "Summary"
    ws3.append(["Generated", dt.datetime.now().strftime("%d %b %Y %H:%M")])
    if partial: ws3.append(["STATUS", partial])
    ws3.append(["Sales orders touched", len(plans)])
    ws3.append(["Packages + shipments to create", sum(len(p["packages"]) for p in plans)])
    ws3.append(["Shipment line rows", len(rows)])
    ws3.append(["Sales returns + receives to create (credit notes)", sum(len(p["returns"]) for p in plans)])
    ws3.append(["Return line rows", len(ret_rows)])
    ws3.append(["Sales orders skipped for human review", len(skipped)])
    ws3.append(["Credit note lines skipped", len(cn_skipped)])
    ws3.append(["API calls used by dry run", Z.calls])
    ws3.append(["Estimated live run calls", sum(len(p["packages"]) for p in plans) * 2 + sum(len(p["returns"]) for p in plans) * 2 + len(plans) * 4])
    ws3.column_dimensions["A"].width = 48; ws3.column_dimensions["B"].width = 20
    ws = wb.create_sheet("Planned shipments")
    hdr(ws, ["SO Number", "Customer", "Source Invoice", "Ship Date (=invoice date)", "Item", "Qty to ship", "Notes"], [14, 40, 16, 20, 45, 12, 50])
    for r in rows: ws.append(r)
    ws2 = wb.create_sheet("Planned returns")
    hdr(ws2, ["SO Number", "Customer", "Credit Note", "Return Date (=CN date)", "Source Invoice", "Item", "Qty back to stock"], [14, 40, 14, 20, 16, 45, 14])
    for r in ret_rows: ws2.append(r)
    ws4 = wb.create_sheet("Skipped, needs human")
    hdr(ws4, ["SO Number", "Customer", "Reason"], [14, 40, 80])
    for r in skipped: ws4.append(r)
    ws5 = wb.create_sheet("Credit note lines skipped")
    hdr(ws5, ["Credit Note", "Customer", "Item", "Qty", "Reason"], [14, 40, 45, 8, 50])
    for r in cn_skipped: ws5.append(r)
    wb.save(f"{OUT}/Backfill_Plan.xlsx")

def _on_term(signum, frame):
    raise SystemExit("run was cancelled (SIGTERM)")

if __name__ == "__main__":
    import signal
    signal.signal(signal.SIGTERM, _on_term)
    main()
